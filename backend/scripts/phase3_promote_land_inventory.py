#!/usr/bin/env python3
"""Preflight and atomically promote the reconciled Phase 3 land inventory."""
from __future__ import annotations

import argparse
import asyncio
import csv
import hashlib
import json
import re
import warnings
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import openpyxl
from sqlalchemy import func, select, text

from app.core.addresses import normalize_address
from app.core.config import settings
from app.core.database import AsyncSessionLocal
from app.models.core import Lot
from app.models.documents import (
    DocType,
    Document,
    DocumentStatus,
    Extraction,
    Ingestion,
    Review,
)
from app.models.land import Agreement, LotTerms
from app.services.ingest import DOCUMENTS_BUCKET, IngestService
from app.services.ocr.extractor import PDFExtractor
from app.services.promotion import PromotionService


ROOT = Path("/home/officehub")
RECONCILIATION_CSV = ROOT / "phase2-agreement-lots.csv"
STAGING_JSON = ROOT / "phase2-staging-results.json"
EXTRACTIONS_JSON = Path("/tmp/phase2_all_land_extractions.json")
WORKBOOK = Path("/tmp/phase2-land-inventory.xlsx")
RAMONA_LPA = Path("/tmp/phase3_evidence/2058922185525.pdf")
RAMONA_ASSIGNMENT = Path("/tmp/phase3_evidence/2058920924670.pdf")
APPROVED_KEYS = {
    "waterside_2024_six",
    "terracon_oak_river_11",
    "forest_grove_phase2",
}
APPROVED_DOCUMENTS = {
    "waterside_2024_six": UUID("d4628f16-9011-4c99-868c-75e745d3223a"),
    "terracon_oak_river_11": UUID("c9eec09e-701a-4a32-86b5-413066a1b1d4"),
    "forest_grove_phase2": UUID("3d9ae465-126d-4e83-867a-4f23335d4f8f"),
}
LEGACY_DOCUMENTS = {
    "morning_glory_12": UUID("376c0803-f458-4f82-a180-d35c5367fb20"),
    "parkview_pointe_10": UUID("f22836fe-6ae6-4ea3-9fc0-cb6c7af3c355"),
    "parkview_pointe_16": UUID("69097875-e4b6-4ea6-8c17-e01727123007"),
}
RAMONA_SHA256 = "71839822808a61c5779756e8074d8bcf16f2e6533fd1a3cd425883f103bb8acd"


AGREEMENT_OVERRIDES: dict[str, dict[str, Any]] = {
    "garson_43": {
        "development_name": "Garson Building Lots",
        "total_purchase_price": "5166700.00",
    },
    "gray_30_35": {
        "vendor_name": "NORTH GRASSIE PROPERTIES INC.",
        "development_name": "Grande Pointe Meadows",
        "municipality": "RM of Ritchot",
        "agreement_date": "2025-03-27",
        "total_purchase_price": "1249765.00",
    },
    "gpm_active_batch": {
        "vendor_name": "NORTH GRASSIE PROPERTIES INC.",
        "development_name": "The Villas at Grande Pointe Meadows",
        "municipality": "RM of Ritchot",
        "agreement_date": "2026-03-31",
        "total_purchase_price": "633970.00",
    },
    "turnberry_812": {
        "agreement_date": "2021-12-22",
        "total_purchase_price": "165000.00",
    },
    "mccallum_6_third_party": {
        "vendor_name": "Mavros Design Build Inc.",
        "development_name": "Wheatland Park",
        "municipality": "RM of Springfield",
        "agreement_date": "2023-06-06",
        "total_purchase_price": "134990.00",
    },
    "gold_sun_34": {"total_purchase_price": "495000.00"},
    "rockall_371": {"total_purchase_price": "465000.00"},
    "a_m_batch": {"total_purchase_price": "355600.00"},
    "stony_16": {"total_purchase_price": "645000.00"},
    "stony_25": {"agreement_date": "2025-02-10", "total_purchase_price": "625000.00"},
    "ramona_154_assignment": {
        "vendor_name": "Waterside Development Corp.",
        "development_name": "Templeton Gates",
        "municipality": "City of Winnipeg",
        "agreement_date": "2025-05-08",
        "total_purchase_price": "165000.00",
        "purchaser_name": "GS Homes Ltd.",
    },
}


SECURITY: dict[str, tuple[str, str, str] | None] = {
    "garson_43": ("3000", "129000", "on closing"),
    "waterside_2024_six": ("3000", "18000", "upon execution"),
    "parkview_pointe_10": ("3000", "30000", "upon execution"),
    "parkview_pointe_16": ("3000", "30000", "upon execution"),
    "gold_sun_34": ("10000", "10000", "on closing"),
    "morning_glory_12": ("2000", "24000", "upon execution / closing"),
    "mccallum_6_third_party": None,
    "champagne_121": None,
    "terracon_oak_river_11": ("3500", "35000", "source agreement cap"),
    "a_m_batch": ("2500", "10000", "landscaping compliance"),
    "gray_30_35": None,
    "buffalo_batch_a": None,
    "buffalo_batch_b": None,
    "gpm_active_batch": ("10000", "60000", "signed clause 12(n) amendment"),
    "buffalo_124_third_party": None,
    "templeton_batch": ("3000", "39000", "upon execution"),
    "ramona_154_assignment": ("3000", "3000", "assumed LPA obligation"),
    "turnberry_812": None,
    "froese_105": None,
    "hall_1273": None,
    "forest_grove_54": ("5000", "25000", "upon execution"),
    "ash_48_52": ("5000", "10000", "upon execution"),
    "rockall_371": ("10000", "10000", "on closing"),
    "forest_grove_phase2": ("5000", "30000", "source agreement"),
    "country_402": ("2500", "2500", "with second deposit / permit trigger"),
    "country_404": ("2500", "2500", "with second deposit / permit trigger"),
    "country_407": ("2500", "2500", "with second deposit / permit trigger"),
    "country_409": ("2500", "2500", "with second deposit / permit trigger"),
    "country_410": ("2500", "2500", "with second deposit / permit trigger"),
    "country_411": ("2500", "2500", "with second deposit / permit trigger"),
    "ella_53_55": ("2500", "2500", "agreement-level security for pair"),
    "ella_37_39": ("2500", "2500", "agreement-level security for pair"),
    "ella_45_47": ("2500", "2500", "agreement-level security for pair"),
    "stony_16": None,
    "stony_25": None,
}


DATE_FORMATS = (
    "%Y-%m-%d",
    "%B %d, %Y",
    "%b %d, %Y",
    "%d %B %Y",
    "%d %b %Y",
)


def clean_json(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, dict):
        return {str(key): clean_json(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean_json(item) for item in value]
    return value


def money(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    match = re.search(r"-?\$?\s*([0-9][0-9,]*(?:\.[0-9]+)?)", str(value))
    if match is None:
        return None
    return Decimal(match.group(1).replace(",", "")).quantize(Decimal("0.01"))


def iso_date(value: Any) -> str | None:
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text_value = str(value or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text_value, fmt).date().isoformat()
        except ValueError:
            pass
    month_only = re.fullmatch(r"([A-Za-z]+)\s+(\d{4})", text_value)
    if month_only:
        return datetime.strptime(f"1 {text_value}", "%d %B %Y").date().isoformat()
    return None


def legal(row: dict[str, str]) -> str:
    block = row["block"].strip()
    lot_number = row["lot"].strip()
    plan = row["plan"].strip()
    if row["agreement_key"] == "mccallum_6_third_party":
        return "TITLE 341829/1"
    if row["agreement_key"] == "rockall_371":
        return "LT 9 PLAN 72953"
    block_text = f"BLK {block} " if block else ""
    return f"{block_text}LT {lot_number} PLAN {plan}".upper()


def workbook_rows() -> dict[int, dict[str, Any]]:
    warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=False)
    sheet = workbook["Original inventory"]
    columns = {
        "sheet_purchase_price": "C",
        "sheet_rebate": "D",
        "sheet_security": "E",
        "contract_deposit_1": "F",
        "contract_deposit_2": "G",
        "contract_deposit_3": "H",
        "tax_adjustment": "I",
        "interest_start": "J",
        "closing": "K",
        "purchase_status": "M",
        "purchase_date": "O",
        "sale_status": "Q",
        "sale_price": "R",
        "lender": "T",
        "paid_1_amount": "U",
        "paid_1_date": "V",
        "paid_2_amount": "W",
        "paid_2_date": "X",
        "paid_3_amount": "Y",
        "paid_3_date": "Z",
        "paid_4_amount": "AA",
        "paid_4_date": "AB",
        "paid_5_amount": "AC",
        "paid_5_date": "AD",
        "paid_6_amount": "AE",
        "paid_6_date": "AF",
        "payout_amount": "AG",
        "payout_date": "AH",
        "calculated_liability": "AJ",
    }
    return {
        row_number: {
            name: clean_json(sheet[f"{column}{row_number}"].value)
            for name, column in columns.items()
        }
        for row_number in range(4, sheet.max_row + 1)
    }


def document_maps() -> tuple[dict[str, UUID], dict[UUID, dict[str, Any]]]:
    key_to_document = {
        item["agreement_key"]: UUID(item["document_id"])
        for item in json.loads(STAGING_JSON.read_text())
    }
    key_to_document.update(LEGACY_DOCUMENTS)
    key_to_document.update(APPROVED_DOCUMENTS)
    extracted = json.loads(EXTRACTIONS_JSON.read_text())
    by_document = {UUID(item["document_id"]): item for item in extracted}
    return key_to_document, by_document


def security_payload(key: str, agreement_lot_count: int) -> dict[str, Any]:
    values = SECURITY[key]
    if values is None:
        return {}
    rate, maximum, trigger = values
    return {
        "rate_per_lot": rate,
        "maximum_amount": maximum,
        "due_trigger": trigger,
        "agreement_lot_count": agreement_lot_count,
    }


def deposit_description(key: str, number: int) -> tuple[str, str | None]:
    if number == 1:
        return "on_signing", "initial contractual deposit"
    if key == "gray_30_35":
        return "milestone", "substantial completion of roads in Phase 6A" if number == 2 else "Hydro work ready for energization in Phase 6A"
    if key in {"buffalo_batch_a", "buffalo_batch_b"}:
        return "milestone", "substantial completion of roads" if number == 2 else "Hydro work ready for energization"
    if key.startswith("country_"):
        return "milestone", "2026-11-01 or building-permit availability, whichever is earlier"
    if key in {"gold_sun_34", "rockall_371"}:
        return "milestone", "after waiver or satisfaction of purchaser conditions"
    return "milestone", f"contractual installment {number}; see operative source"


def build_payloads() -> tuple[dict[str, dict[str, Any]], dict[str, UUID]]:
    with RECONCILIATION_CSV.open(newline="", encoding="utf-8-sig") as handle:
        all_rows = list(csv.DictReader(handle))
    active = [row for row in all_rows if row["disposition"] == "active_match"]
    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    all_grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in active:
        grouped[row["agreement_key"]].append(row)
    for row in all_rows:
        all_grouped[row["agreement_key"]].append(row)

    sheet_rows = workbook_rows()
    key_to_document, extraction_by_document = document_maps()
    payloads: dict[str, dict[str, Any]] = {}
    for key, rows in grouped.items():
        representative = rows[0]
        source_record = extraction_by_document.get(key_to_document.get(key))
        extracted = deepcopy((source_record or {}).get("extracted_payload") or {})
        extracted_agreement = extracted.get("agreement", {})
        if not isinstance(extracted_agreement, dict):
            extracted_agreement = {}
        agreement = dict(extracted_agreement)
        agreement.update(AGREEMENT_OVERRIDES.get(key, {}))
        agreement["agreement_date"] = iso_date(agreement.get("agreement_date"))
        agreement["vendor_name"] = agreement.get("vendor_name") or representative["vendor_block"].split(" — ", 1)[0]
        agreement["development_name"] = agreement.get("development_name") or representative["agreement_title"]
        agreement["municipality"] = agreement.get("municipality") or "Manitoba"
        parsed_total = money(agreement.get("total_purchase_price"))
        agreement["total_purchase_price"] = str(parsed_total) if parsed_total is not None else None
        agreement["vendor_name_source_text"] = representative["vendor_block"]

        lot_payloads: list[dict[str, Any]] = []
        seen_sheet_rows: Counter[int] = Counter()
        for row in rows:
            sheet_row_number = int(row["active_sheet_row"])
            sheet = sheet_rows[sheet_row_number]
            seen_sheet_rows[sheet_row_number] += 1
            pair_member = len([item for item in rows if item["active_sheet_row"] == row["active_sheet_row"]]) > 1
            primary_pair_member = not pair_member or seen_sheet_rows[sheet_row_number] == 1
            source_price = money(row["purchase_price"])
            if key == "gold_sun_34":
                source_price = Decimal("495000.00")
            elif key == "rockall_371":
                source_price = Decimal("465000.00")
            elif key == "garson_43" or pair_member:
                source_price = None

            metadata = {
                "sheet_row": sheet_row_number,
                "sheet_purchase_price": sheet["sheet_purchase_price"],
                "sheet_status": {
                    "purchase": sheet["purchase_status"],
                    "sale": sheet["sale_status"],
                },
                "source_purchase_price_text": row["purchase_price"],
                "tax_terms": "purchase price excludes GST; applicable GST is a separate tax obligation" if "GST" in row["purchase_price"].upper() or key == "garson_43" else None,
                "rebate_terms": row["rebate"],
                "security_terms": row["security_deposit"],
                "damage_deposit": "2500.00" if key.startswith("country_") or (key.startswith("ella_") and primary_pair_member) else None,
                "date_terms_source": row["date_terms"],
                "sheet_dates": {
                    "tax_adjustment": sheet["tax_adjustment"],
                    "interest_start": sheet["interest_start"],
                    "closing": sheet["closing"],
                    "purchase_date": sheet["purchase_date"],
                },
                "payment_history_source": "original inventory workbook U:AF; blanks preserved",
                "payment_history": [
                    {
                        "number": number,
                        "amount": sheet[f"paid_{number}_amount"],
                        "date": sheet[f"paid_{number}_date"],
                    }
                    for number in range(1, 7)
                    if sheet[f"paid_{number}_amount"] not in (None, "") or sheet[f"paid_{number}_date"] not in (None, "")
                ],
                "payout": {
                    "amount": sheet["payout_amount"],
                    "date": sheet["payout_date"],
                    "reliability": "imported_as_is; no reliable OTP-to-payout reconciliation convention",
                },
                "calculated_liability_from_sheet": sheet["calculated_liability"],
                "lender_annotation": sheet["lender"],
                "source_notes": row["notes"],
            }
            if key == "garson_43":
                metadata["price_allocation"] = "agreement_only"
            if pair_member:
                metadata["price_allocation"] = "agreement_total_unallocated"
                metadata["agreement_pair_total"] = str(money(row["purchase_price"]))
                metadata["combined_sheet_row"] = True
                metadata["agreement_level_money_holder"] = primary_pair_member
                if not primary_pair_member:
                    metadata["payment_history"] = []
                    metadata["payout"] = {
                        "amount": None,
                        "date": None,
                        "reliability": "agreement-level values are stored on the pair's primary physical lot",
                    }
                    metadata["calculated_liability_from_sheet"] = None
            if key in {"mccallum_6_third_party", "forest_grove_54"}:
                metadata["payout_reconciliation_flag"] = "M8: retained without recomputation"
            if key == "mccallum_6_third_party":
                metadata["legal_description_reliability"] = "scan-derived title; verification remains open"

            schedule: list[dict[str, Any]] = []
            if primary_pair_member:
                amounts = [
                    money(sheet["contract_deposit_1"]),
                    money(sheet["contract_deposit_2"]),
                    money(sheet["contract_deposit_3"]),
                ]
                if key in {"gold_sun_34", "rockall_371"}:
                    amounts[1] = Decimal("25000.00")
                if key.startswith("ella_"):
                    amounts = [Decimal("18400.00"), None, None]
                for number, amount in enumerate(amounts, start=1):
                    if amount is None or amount == 0:
                        continue
                    trigger_type, description = deposit_description(key, number)
                    schedule.append(
                        {
                            "deposit_number": number,
                            "amount": str(amount),
                            "due_date": agreement["agreement_date"] if number == 1 else None,
                            "trigger_type": trigger_type,
                            "trigger_description": description,
                            "source_text": row["deposit_schedule"],
                            "paid_amount": sheet[f"paid_{number}_amount"],
                            "paid_at": sheet[f"paid_{number}_date"],
                        }
                    )

            lot_payloads.append(
                {
                    "legal_description_normalized": legal(row),
                    "legal_description_raw": row["plan"] if key == "mccallum_6_third_party" else legal(row),
                    "lot_number": row["lot"] or None,
                    "block": row["block"] or None,
                    "plan": row["plan"] or None,
                    "civic_address": row["civic_address"],
                    "purchase_price": str(source_price) if source_price is not None else None,
                    "lot_notes": row["notes"] or None,
                    "deposit_schedule": schedule,
                    "sale_type": "presale" if str(sheet["sale_status"] or "").strip().lower() == "sold" else None,
                    "metadata": clean_json(metadata),
                }
            )

        source_documents = [
            {
                "name": representative["operative_document_set"],
                "sha1": representative["source_sha1"],
                "document_id": str(key_to_document[key]) if key in key_to_document else None,
                "role": "operative source set",
            }
        ]
        if key == "ramona_154_assignment":
            source_documents = [
                {
                    "name": RAMONA_LPA.name,
                    "box_file_id": "2058922185525",
                    "sha1": "dab817c76893506345e974b45ea38e7b42a531c4",
                    "sha256": RAMONA_SHA256,
                    "role": "operative signed Waterside/GS Homes LPA",
                },
                {
                    "name": RAMONA_ASSIGNMENT.name,
                    "box_file_id": "2058920924670",
                    "sha1": "b01f25b1bf3701f3d5d4b465e85a6290ec3c9475",
                    "role": "unsigned assignment; chain evidence only",
                },
            ]
        open_items: list[str] = []
        if key == "ramona_154_assignment":
            open_items.append("No executed GS Homes-to-Connection Homes assignment counterpart was found; non-blocking legal follow-up remains open.")
        if key == "mccallum_6_third_party":
            open_items.append("Title 341829/1 is scan-derived and remains flagged for verification.")
            open_items.append("The buyer signed the offer on 2023-06-06; the seller-acceptance signature block in the located copy is blank.")
        if key == "garson_43":
            open_items.append("Rows 24 and 25 both use 32 Boulder Crescent for distinct legal lots; civic correction remains open.")

        payloads[key] = {
            "agreement_key": key,
            "agreement": clean_json(agreement),
            "lots": lot_payloads,
            "security_deposit": security_payload(key, len(all_grouped[key])),
            "development_guidelines": extracted.get("development_guidelines", {}),
            "notable_clauses": extracted.get("notable_clauses", []),
            "source_documents": source_documents,
            "operative_selection": {
                "rationale": representative["selection_reason"],
                "operative_set": representative["operative_document_set"],
            },
            "field_provenance": {
                "legal_and_lot_terms": "phase2-agreement-lots.csv, source-reviewed against operative agreement set",
                "payment_and_payout_history": "Original inventory workbook; existing values imported as-is",
                "money_decisions": "Accepted M1-M8 decision register",
            },
            "source_terms": {
                "agreement_title": representative["agreement_title"],
                "vendor_block": representative["vendor_block"],
                "date_terms": representative["date_terms"],
                "security_terms": sorted({row["security_deposit"] for row in rows}),
                "rebate_terms": sorted({row["rebate"] for row in rows}),
                "deposit_terms": sorted({row["deposit_schedule"] for row in rows}),
                "agreement_lot_count": len(all_grouped[key]),
                "active_physical_lot_count": len(rows),
            },
            "open_legal_items": open_items,
            "reconciled_agreement_fields": ["agreement_date"] if key == "mccallum_6_third_party" else [],
        }
    return payloads, key_to_document


def validate_payloads(payloads: dict[str, dict[str, Any]]) -> dict[str, Any]:
    errors: list[str] = []
    legal_rows: list[tuple[str, str]] = []
    sheet_rows: set[int] = set()
    presales: list[str] = []
    for key, payload in payloads.items():
        agreement = payload["agreement"]
        for field in ("vendor_name", "development_name", "municipality", "agreement_date", "total_purchase_price"):
            if agreement.get(field) in (None, ""):
                errors.append(f"{key}: agreement.{field} is missing")
        for lot in payload["lots"]:
            legal_rows.append((lot["legal_description_normalized"], key))
            sheet_rows.add(int(lot["metadata"]["sheet_row"]))
            if lot["sale_type"] == "presale":
                presales.append(lot["civic_address"])
            if lot["purchase_price"] is None and lot["metadata"].get("price_allocation") not in {"agreement_only", "agreement_total_unallocated"}:
                errors.append(f"{key}: {lot['legal_description_normalized']} has no price allocation")
    duplicates = [legal for legal, count in Counter(item[0] for item in legal_rows).items() if count > 1]
    if duplicates:
        errors.append(f"duplicate legal descriptions: {duplicates}")
    if len(payloads) != 35:
        errors.append(f"expected 35 agreement sets, got {len(payloads)}")
    if len(legal_rows) != 138:
        errors.append(f"expected 138 physical active lots, got {len(legal_rows)}")
    if len(sheet_rows) != 135:
        errors.append(f"expected 135 active sheet rows, got {len(sheet_rows)}")
    if len(presales) != 16:
        errors.append(f"expected 16 sold/presale rows, got {len(presales)}")
    return {
        "errors": errors,
        "agreement_sets": len(payloads),
        "physical_lots": len(legal_rows),
        "sheet_rows": len(sheet_rows),
        "presales": sorted(presales),
    }


async def resolve_database_state(
    payloads: dict[str, dict[str, Any]], key_to_document: dict[str, UUID]
) -> dict[str, Any]:
    async with AsyncSessionLocal() as db:
        current_revision = await db.scalar(text("SELECT version_num FROM alembic_version"))
        documents: dict[str, Any] = {}
        for key, document_id in key_to_document.items():
            document = await db.get(Document, document_id)
            promoted = await db.scalar(select(Agreement.id).where(Agreement.document_id == document_id))
            documents[key] = {
                "document_id": str(document_id),
                "status": document.status.value if document else None,
                "promoted_agreement_id": str(promoted) if promoted else None,
            }

        planned: list[dict[str, Any]] = []
        service = PromotionService(db)
        for key, payload in payloads.items():
            for lot_payload in payload["lots"]:
                legal_description = lot_payload["legal_description_normalized"]
                existing = await db.scalar(select(Lot).where(Lot.legal_description_normalized == legal_description))
                candidates = await service._property_candidates(lot_payload["civic_address"])
                planned.append(
                    {
                        "agreement_key": key,
                        "legal_description": legal_description,
                        "civic_address": lot_payload["civic_address"],
                        "lot_action": "match" if existing else "create",
                        "existing_lot_id": str(existing.id) if existing else None,
                        "current_property_id": str(existing.property_id) if existing and existing.property_id else None,
                        "property_candidates": [
                            {"id": str(candidate.id), "address": candidate.address}
                            for candidate in candidates
                        ],
                    }
                )
        return {
            "database_revision": current_revision,
            "documents": documents,
            "planned_lots": planned,
            "planned_lot_actions": dict(Counter(item["lot_action"] for item in planned)),
            "planned_property_candidate_counts": dict(Counter(len(item["property_candidates"]) for item in planned)),
        }


async def stage_ramona(db: Any, payload: dict[str, Any]) -> tuple[Document, Extraction]:
    existing = await db.scalar(select(Document).where(Document.checksum_sha256 == RAMONA_SHA256))
    if existing is not None:
        extraction = await db.scalar(
            select(Extraction)
            .join(Ingestion, Extraction.ingestion_id == Ingestion.id)
            .where(Ingestion.document_id == existing.id)
            .order_by(Extraction.created_at.desc())
            .limit(1)
        )
        if extraction is None:
            raise ValueError("The staged Ramona LPA has no extraction")
        return existing, extraction

    file_bytes = RAMONA_LPA.read_bytes()
    if hashlib.sha256(file_bytes).hexdigest() != RAMONA_SHA256:
        raise ValueError("Ramona LPA checksum changed")
    org_id = UUID(str(settings.default_org_id))
    object_suffix = "-154-Ramona-signed-Waterside-GS-Homes-LPA.pdf"
    ingest_service = IngestService(db)
    client = ingest_service._s3_client()
    existing_objects = client.list_objects_v2(
        Bucket=DOCUMENTS_BUCKET,
        Prefix="inbox/",
    ).get("Contents", [])
    matching_objects = [
        item for item in existing_objects if item.get("Key", "").endswith(object_suffix)
    ]
    if matching_objects:
        minio_key = max(
            matching_objects,
            key=lambda item: item.get("LastModified") or datetime.min.replace(tzinfo=timezone.utc),
        )["Key"]
    else:
        minio_key = f"inbox/{uuid4()}{object_suffix}"
        ingest_service._upload_pdf(temp_path=RAMONA_LPA, minio_key=minio_key)
    ocr_started = datetime.now(timezone.utc)
    ocr = await asyncio.to_thread(PDFExtractor().extract, RAMONA_LPA)
    now = datetime.now(timezone.utc)
    document = Document(
        org_id=org_id,
        doc_type=DocType.LAND_OTP,
        status=DocumentStatus.IN_REVIEW,
        original_filename="154 Ramona - signed Waterside GS Homes LPA.pdf",
        minio_bucket=DOCUMENTS_BUCKET,
        minio_key=minio_key,
        file_size_bytes=len(file_bytes),
        checksum_sha256=RAMONA_SHA256,
    )
    db.add(document)
    await db.flush()
    ingestion = Ingestion(
        document_id=document.id,
        ocr_method="manual",
        ocr_text=ocr.raw_text,
        ocr_confidence=Decimal("1.000"),
        page_count=ocr.total_pages,
        started_at=ocr_started,
        completed_at=now,
    )
    db.add(ingestion)
    await db.flush()
    extraction = Extraction(
        ingestion_id=ingestion.id,
        model_provider="manual_source_review",
        model_version="phase3-v1",
        prompt_version="land-otp-v8",
        extracted_payload=payload,
        field_confidences={},
        low_confidence_fields=[],
    )
    db.add(extraction)
    await db.flush()
    return document, extraction


async def latest_extraction(db: Any, document_id: UUID) -> Extraction:
    extraction = await db.scalar(
        select(Extraction)
        .join(Ingestion, Extraction.ingestion_id == Ingestion.id)
        .where(Ingestion.document_id == document_id)
        .order_by(Extraction.created_at.desc())
        .limit(1)
    )
    if extraction is None:
        raise ValueError(f"No extraction for {document_id}")
    return extraction


async def latest_or_manual_extraction(
    db: Any,
    document_id: UUID,
    payload: dict[str, Any],
) -> Extraction:
    extraction = await db.scalar(
        select(Extraction)
        .join(Ingestion, Extraction.ingestion_id == Ingestion.id)
        .where(Ingestion.document_id == document_id)
        .order_by(Extraction.created_at.desc())
        .limit(1)
    )
    if extraction is not None:
        return extraction
    ingestion = await db.scalar(
        select(Ingestion)
        .where(Ingestion.document_id == document_id)
        .order_by(Ingestion.completed_at.desc().nullslast())
        .limit(1)
    )
    if ingestion is None:
        raise ValueError(f"No ingestion for {document_id}")
    extraction = Extraction(
        ingestion_id=ingestion.id,
        model_provider="manual_source_review",
        model_version="phase3-v1",
        prompt_version="land-otp-v8",
        extracted_payload=payload,
        field_confidences={},
        low_confidence_fields=[],
    )
    db.add(extraction)
    await db.flush()
    return extraction


async def collect_link_results(db: Any, payloads: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    service = PromotionService(db)
    for key, payload in payloads.items():
        for lot_payload in payload["lots"]:
            lot = await db.scalar(select(Lot).where(Lot.legal_description_normalized == lot_payload["legal_description_normalized"]))
            if lot is None:
                results.append({"agreement_key": key, "legal_description": lot_payload["legal_description_normalized"], "civic_address": lot_payload["civic_address"], "status": "lot_missing"})
                continue
            candidates = await service._property_candidates(lot_payload["civic_address"])
            if lot.property_id is not None and any(
                candidate.id == lot.property_id for candidate in candidates
            ):
                status = "linked"
            elif lot.property_id is not None and candidates:
                status = "existing_property_mismatch"
            elif lot.property_id is not None:
                status = "existing_link_not_civic_verified"
            elif not candidates:
                status = "no_property_match"
            elif len(candidates) > 1:
                status = "ambiguous_property_match"
            else:
                conflicting = await db.scalar(select(Lot.id).where(Lot.property_id == candidates[0].id, Lot.id != lot.id).limit(1))
                status = "property_link_conflict" if conflicting else "not_linked"
            results.append(
                {
                    "agreement_key": key,
                    "lot_id": str(lot.id),
                    "legal_description": lot.legal_description_normalized,
                    "civic_address": lot_payload["civic_address"],
                    "status": status,
                    "property_id": str(lot.property_id) if lot.property_id else None,
                    "candidates": [{"id": str(item.id), "address": item.address} for item in candidates],
                }
            )
    return results


async def execute(payloads: dict[str, dict[str, Any]], key_to_document: dict[str, UUID]) -> dict[str, Any]:
    async with AsyncSessionLocal() as db:
        try:
            await db.execute(text("SET TRANSACTION ISOLATION LEVEL SERIALIZABLE"))
            revision = await db.scalar(text("SELECT version_num FROM alembic_version"))
            if revision != "20260914_0044":
                raise ValueError(f"Expected migration 20260914_0044, found {revision}")
            service = PromotionService(db)
            promoted: list[dict[str, Any]] = []
            annotated: list[dict[str, Any]] = []

            for key in sorted(APPROVED_KEYS):
                agreement = await db.scalar(select(Agreement).where(Agreement.document_id == APPROVED_DOCUMENTS[key]))
                if agreement is None:
                    raise ValueError(f"Approved agreement missing for {key}")
                lot_ids = await service.annotate_promoted_land_agreement(agreement.id, payloads[key])
                annotated.append({"agreement_key": key, "agreement_id": str(agreement.id), "lots": len(lot_ids)})

            for key in sorted(set(payloads) - APPROVED_KEYS):
                if key == "ramona_154_assignment":
                    ramona_document, extraction = await stage_ramona(db, payloads[key])
                    key_to_document[key] = ramona_document.id
                document_id = key_to_document[key]
                document = await db.get(Document, document_id)
                if document is None:
                    raise ValueError(f"Document missing for {key}: {document_id}")
                if await db.scalar(select(Agreement.id).where(Agreement.document_id == document_id)) is not None:
                    raise ValueError(f"Document unexpectedly already promoted for {key}")
                if key != "ramona_154_assignment":
                    extraction = await latest_or_manual_extraction(
                        db,
                        document_id,
                        payloads[key],
                    )
                review = Review(
                    extraction_id=extraction.id,
                    reviewed_payload=payloads[key],
                    edited_fields=["agreement", "lots", "security_deposit", "source_documents", "operative_selection", "field_provenance"],
                    decision="approved",
                    reviewed_at=datetime.now(timezone.utc),
                )
                db.add(review)
                await db.flush()
                result = await service.promote(review.id, commit=False)
                promoted.append(
                    {
                        "agreement_key": key,
                        "document_id": str(result.document_id),
                        "review_id": str(result.review_id),
                        "agreement_id": str(result.agreement_id),
                        "lots_created": result.lots_created,
                        "lots_matched": result.lots_matched,
                    }
                )

            await db.flush()
            link_results = await collect_link_results(db, payloads)
            counts = {
                "agreements": await db.scalar(select(func.count()).select_from(Agreement)),
                "lots": await db.scalar(select(func.count()).select_from(Lot)),
                "lot_terms": await db.scalar(select(func.count()).select_from(LotTerms)),
            }
            await db.commit()
            return {
                "status": "committed",
                "committed_at": datetime.now(timezone.utc).isoformat(),
                "promoted": promoted,
                "annotated_existing": annotated,
                "totals_after_commit": counts,
                "property_links": link_results,
                "property_link_summary": dict(Counter(item["status"] for item in link_results)),
                "failed_property_links": [item for item in link_results if item["status"] != "linked"],
            }
        except Exception:
            await db.rollback()
            raise


async def async_main(args: argparse.Namespace) -> dict[str, Any]:
    payloads, key_to_document = build_payloads()
    validation = validate_payloads(payloads)
    result: dict[str, Any] = {"validation": validation}
    if validation["errors"]:
        return result
    if args.mode in {"preflight", "execute"}:
        result["database"] = await resolve_database_state(payloads, key_to_document)
    if args.mode == "execute":
        result["execution"] = await execute(payloads, key_to_document)
    return clean_json(result)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=("prepare", "preflight", "execute"))
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = asyncio.run(async_main(args))
    serialized = json.dumps(result, indent=2, sort_keys=True)
    if args.output:
        args.output.write_text(serialized + "\n")
    print(serialized)
    if result["validation"]["errors"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
